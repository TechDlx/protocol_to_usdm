"""Stage B: write the USDM Excel workbook from a confirmed review. Deterministic; no model involved.

The workbook is the usdm4-excel legacy single-workbook dialect (docs/usdm_workbook_spec.md, D1),
laid out exactly as usdm4-excel's importer reads it:

- ordinary sheets: headers on row 1, one record per row, columns in layout order;
- `study`: the key/value block, a blank row, then the governance dates table;
- `studyDesign`: the key/value block (with mainTimeline/otherTimelines naming the timeline sheets),
  a blank row, then the arm-by-epoch grid of element names;
- one sheet per timeline: Name/Label/Description/Condition in columns A-B, the timepoint heading
  rows (name, description, label, type, default, condition, epoch, encounter, timeline) in column C
  with one column per timepoint, then from row 10 the activity table
  (`Parent Activity | Child Activity | BC/Procedure/Timeline | X marks`).

Cell text: a controlled-terminology cell is written as the preferred term of its resolution (the
importer matches preferred terms, not synonyms), keeping an "Other=<reason>"; dates are written as
text in the `yyyy-mm-dd 00:00:00` form the importer parses; everything else as reviewed.
"""

import html
import re
import zipfile
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from backend.models.extraction import ExtractedField, ExtractionSheets, TerminologyStatus
from backend.pipeline.terminology.ct import MULTI_SEPARATOR
from backend.pipeline.workbook.cells import OTHER_PREFIX
from backend.pipeline.workbook.formats import ValueFormat
from backend.pipeline.workbook.layout import (
    DATES,
    EXIT,
    SCHEDULE,
    SHEETS,
    STUDY,
    STUDY_CELLS,
    STUDY_DESIGN,
    TIMELINES,
    TIMEPOINTS,
    ColumnSpec,
    SheetKind,
    SheetSpec,
)
from backend.pipeline.workbook.sources import sheet_rows

HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
HEADER_FONT = Font(bold=True)
KEY_COLUMN_WIDTH = 28.0
VALUE_COLUMN_WIDTH = 40.0
TABLE_COLUMN_WIDTH = 24.0
TIMEPOINT_COLUMN_WIDTH = 16.0
TIMELINE_HEADINGS = (
    "name",
    "description",
    "label",
    "type",
    "default",
    "condition",
    "epoch",
    "encounter",
    "timeline",
)
ACTIVITY_HEADER = ("Parent Activity", "Child Activity", "BC/Procedure/Timeline")
ACTIVITY_HEADER_ROW = len(TIMELINE_HEADINGS) + 1  # row 10
GRID_CORNER = "Epoch/Arms"
CREATOR = "protocol_to_usdm"
FIXED_TIMESTAMP = datetime(2000, 1, 1, tzinfo=UTC)
MARK = "X"

# Review views written through special layouts rather than as sheets of their own.
SPECIAL = {
    STUDY.key,
    DATES.key,
    STUDY_DESIGN.key,
    TIMELINES.key,
    TIMEPOINTS.key,
    SCHEDULE.key,
    STUDY_CELLS.key,
}


@dataclass
class WorkbookSummary:
    sheets: dict[str, int] = field(default_factory=dict)  # workbook sheet -> data rows written
    warnings: list[str] = field(default_factory=list)


def cell_text(column: ColumnSpec, value: ExtractedField[str] | None) -> str:
    """The text written for one reviewed value."""
    if value is None or value.value is None:
        return ""
    text = value.value.strip()
    if not text:
        return ""
    t = value.terminology
    if (
        column.ct is not None
        and t is not None
        and t.status == TerminologyStatus.EXACT
        and t.preferred_term
    ):
        terms = [p.strip() for p in t.preferred_term.split(", ")]
        items = (
            [i.strip() for i in text.split(MULTI_SEPARATOR) if i.strip()]
            if column.multi
            else [text]
        )
        if len(terms) == len(items):
            # The amendment reason reader splits on "," without trimming, so no spaces there.
            separator = "," if column.other_allowed else ", "
            return separator.join(
                _with_other(column, item, term) for item, term in zip(items, terms, strict=True)
            )
    if column.format == ValueFormat.DATE:
        return f"{text} 00:00:00"
    if column.xhtml:
        return f"<p>{html.escape(text, quote=False)}</p>"
    return text


def _with_other(column: ColumnSpec, item: str, term: str) -> str:
    if column.other_allowed and item.casefold().startswith(OTHER_PREFIX) and "=" in item:
        return f"{term}={item.split('=', 1)[1].strip()}"
    return term


def _value(record: Any, column: ColumnSpec) -> str:
    return cell_text(column, getattr(record, column.field)) if column.field else ""


def _heading(timepoint: Any, heading: str) -> str:
    """One timepoint's value for a heading row of a timeline sheet."""
    if heading == "timeline":
        return ""  # sub-timeline references are not modelled yet
    if heading == "type":
        return timepoint.type.value or "Activity"
    value: str = getattr(timepoint, heading).value or ""
    if heading == "default" and value == EXIT:
        return EXIT.upper()  # the importer's exit marker
    return value


def _bold(ws: Worksheet, row: int, first_column: int, last_column: int) -> None:
    for col in range(first_column, last_column + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL


def _widths(ws: Worksheet, widths: dict[int, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


class WorkbookWriter:
    def __init__(self, sheets: ExtractionSheets) -> None:
        self.sheets = sheets
        self.book = Workbook()
        self.summary = WorkbookSummary()
        self._first = True

    def _sheet(self, name: str) -> Worksheet:
        if self._first:
            ws = self.book.active
            assert ws is not None
            ws.title = name
            self._first = False
            return ws
        return self.book.create_sheet(name)

    # ----- ordinary sheets ------------------------------------------------------------------

    def _key_value_block(self, ws: Worksheet, spec: SheetSpec, extra: list[tuple[str, str]]) -> int:
        record = next(iter(sheet_rows(self.sheets, spec)), None)
        rows = [(c.header, _value(record, c) if record is not None else "") for c in spec.columns]
        rows += extra
        for r, (key, text) in enumerate(rows, start=1):
            ws.cell(row=r, column=1, value=key)
            ws.cell(row=r, column=2, value=text)
            ws.cell(row=r, column=1).font = HEADER_FONT
            ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        _widths(ws, {1: KEY_COLUMN_WIDTH, 2: VALUE_COLUMN_WIDTH})
        return len(rows)

    def _table(self, ws: Worksheet, spec: SheetSpec, first_row: int = 1) -> int:
        header_row = first_row
        for c, column in enumerate(spec.columns, start=1):
            ws.cell(row=header_row, column=c, value=column.header)
        _bold(ws, header_row, 1, len(spec.columns))
        rows = sheet_rows(self.sheets, spec)
        for r, record in enumerate(rows, start=header_row + 1):
            for c, column in enumerate(spec.columns, start=1):
                cell = ws.cell(row=r, column=c, value=_value(record, column))
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        _widths(ws, {c: TABLE_COLUMN_WIDTH for c in range(1, len(spec.columns) + 1)})
        return len(rows)

    def _write_table_sheet(self, spec: SheetSpec) -> None:
        ws = self._sheet(spec.workbook_sheet)
        ws.freeze_panes = "A2"
        self.summary.sheets[spec.workbook_sheet] = self._table(ws, spec)

    # ----- special layouts ------------------------------------------------------------------

    def _write_study(self) -> None:
        ws = self._sheet(STUDY.workbook_sheet)
        block = self._key_value_block(ws, STUDY, [])
        # A blank key row ends the block; the dates table header follows on the next row.
        dates = self._table(ws, DATES, first_row=block + 2)
        self.summary.sheets[STUDY.workbook_sheet] = block + dates

    def _timelines(self) -> list[Any]:
        return sheet_rows(self.sheets, TIMELINES)

    def _write_study_design(self) -> None:
        ws = self._sheet(STUDY_DESIGN.workbook_sheet)
        timelines = self._timelines()
        main = [t for t in timelines if (t.main.value or "").upper() in ("Y", "YES", "TRUE")]
        others = [t for t in timelines if t not in main]
        extra = [
            ("mainTimeline", (main[0].sheet_name.value or "") if main else ""),
            ("otherTimelines", ", ".join(t.sheet_name.value or "" for t in others)),
        ]
        # The layout lists mainTimeline/otherTimelines as not extracted; fill them here instead.
        spec_rows = [
            c for c in STUDY_DESIGN.columns if c.header not in ("mainTimeline", "otherTimelines")
        ]
        record = next(iter(sheet_rows(self.sheets, STUDY_DESIGN)), None)
        rows = [
            (c.header, _value(record, c) if record is not None else "") for c in spec_rows
        ] + extra
        for r, (key, text) in enumerate(rows, start=1):
            ws.cell(row=r, column=1, value=key).font = HEADER_FONT
            ws.cell(row=r, column=2, value=text).alignment = Alignment(
                wrap_text=True, vertical="top"
            )

        grid_row = len(rows) + 2
        epochs = [
            e.name.value or ""
            for e in (self.sheets.schedule.epochs if self.sheets.schedule else [])
        ]
        cells: dict[tuple[str, str], str] = {}
        arms: list[str] = []
        for cell in sheet_rows(self.sheets, STUDY_CELLS):
            arm, epoch = cell.arm.value or "", cell.epoch.value or ""
            if arm not in arms:
                arms.append(arm)
            cells[(arm, epoch)] = cell.elements.value or ""
        if epochs and arms:
            ws.cell(row=grid_row, column=1, value=GRID_CORNER)
            for c, epoch in enumerate(epochs, start=2):
                ws.cell(row=grid_row, column=c, value=epoch)
            _bold(ws, grid_row, 1, len(epochs) + 1)
            for r, arm in enumerate(arms, start=grid_row + 1):
                ws.cell(row=r, column=1, value=arm).font = HEADER_FONT
                for c, epoch in enumerate(epochs, start=2):
                    ws.cell(row=r, column=c, value=cells.get((arm, epoch), ""))
        elif epochs or arms:
            self.summary.warnings.append(
                "studyDesign: the arm-by-epoch grid needs both arms and epochs; it was left out"
            )
        _widths(
            ws,
            {
                1: KEY_COLUMN_WIDTH,
                2: VALUE_COLUMN_WIDTH,
                **{c: TABLE_COLUMN_WIDTH for c in range(3, len(epochs) + 2)},
            },
        )
        self.summary.sheets[STUDY_DESIGN.workbook_sheet] = len(rows) + len(arms)

    def _write_timelines(self) -> None:
        timepoints: dict[str, list[Any]] = defaultdict(list)
        for tp in sheet_rows(self.sheets, TIMEPOINTS):
            timepoints[tp.timeline.value or ""].append(tp)
        rows_by_timeline: dict[str, list[Any]] = defaultdict(list)
        for row in sheet_rows(self.sheets, SCHEDULE):
            rows_by_timeline[row.timeline.value or ""].append(row)
        bc_column = SCHEDULE.column("biomedical_concepts")

        for timeline in self._timelines():
            name = timeline.name.value or ""
            sheet_name = timeline.sheet_name.value or name
            ws = self._sheet(sheet_name)
            meta = [
                ("Name", name),
                ("Label", timeline.label.value or ""),
                ("Description", timeline.description.value or ""),
                ("Condition", timeline.entry_condition.value or ""),
            ]
            for r, (key, text) in enumerate(meta, start=1):
                ws.cell(row=r, column=1, value=key).font = HEADER_FONT
                ws.cell(row=r, column=2, value=text).alignment = Alignment(
                    wrap_text=True, vertical="top"
                )

            columns = timepoints.get(name, [])
            for r, heading in enumerate(TIMELINE_HEADINGS, start=1):
                ws.cell(row=r, column=3, value=heading).font = HEADER_FONT
                for c, tp in enumerate(columns, start=4):
                    ws.cell(row=r, column=c, value=_heading(tp, heading)).alignment = Alignment(
                        horizontal="center", wrap_text=True
                    )

            for c, text in enumerate(ACTIVITY_HEADER, start=1):
                ws.cell(row=ACTIVITY_HEADER_ROW, column=c, value=text)
            _bold(ws, ACTIVITY_HEADER_ROW, 1, len(ACTIVITY_HEADER))
            names = [tp.name.value or "" for tp in columns]
            written = 0
            for r, row in enumerate(rows_by_timeline.get(name, []), start=ACTIVITY_HEADER_ROW + 1):
                marks = {m.strip() for m in (row.scheduled_at.value or "").split(MULTI_SEPARATOR)}
                concepts = [
                    b.strip()
                    for b in cell_text(bc_column, row.biomedical_concepts).split(MULTI_SEPARATOR)
                    if b.strip()
                ]
                ws.cell(row=r, column=2, value=row.activity.value or "")
                ws.cell(row=r, column=3, value=", ".join(f"BC: {b}" for b in concepts))
                for c, tp_name in enumerate(names, start=4):
                    if tp_name in marks:
                        ws.cell(row=r, column=c, value=MARK).alignment = Alignment(
                            horizontal="center"
                        )
                unknown = marks - set(names) - {""}
                if unknown:
                    self.summary.warnings.append(
                        f"{sheet_name}: '{row.activity.value}' is scheduled at timepoints not "
                        f"on this timeline ({', '.join(sorted(unknown))}); marks not written"
                    )
                written += 1
            _widths(
                ws,
                {
                    1: KEY_COLUMN_WIDTH,
                    2: KEY_COLUMN_WIDTH,
                    3: KEY_COLUMN_WIDTH,
                    **{c: TIMEPOINT_COLUMN_WIDTH for c in range(4, len(names) + 4)},
                },
            )
            ws.freeze_panes = ws.cell(row=ACTIVITY_HEADER_ROW + 1, column=4)
            self.summary.sheets[sheet_name] = written

    # ----- whole workbook -------------------------------------------------------------------

    def write(self, path: Path, timestamp: datetime | None = None) -> WorkbookSummary:
        """Write the workbook. With a timestamp (the review's confirmation time) the file is
        byte-identical for identical content: document properties and zip entry times are pinned."""
        self._write_study()
        self._write_study_design()
        for spec in SHEETS.values():
            if spec.key in SPECIAL or spec.kind == SheetKind.KEY_VALUE:
                continue
            self._write_table_sheet(spec)
        self._write_timelines()
        path.parent.mkdir(parents=True, exist_ok=True)
        when = (timestamp or FIXED_TIMESTAMP).astimezone(UTC).replace(tzinfo=None, microsecond=0)
        props = self.book.properties
        props.creator = CREATOR
        props.lastModifiedBy = CREATOR
        props.created = when
        props.modified = when
        tmp = path.with_suffix(".tmp.xlsx")
        self.book.save(tmp)
        _pin_zip_times(tmp, when)
        tmp.replace(path)
        return self.summary


_MODIFIED = re.compile(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)")


def _pin_zip_times(path: Path, when: datetime) -> None:
    """Rewrite the xlsx zip with every entry dated `when`, and the document's modified time set to
    it too: openpyxl stamps the save time in both places."""
    stamp = (max(when.year, 1980), when.month, when.day, when.hour, when.minute, when.second)
    modified = when.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    rewritten = path.with_suffix(".pinned")
    with (
        zipfile.ZipFile(path) as source,
        zipfile.ZipFile(rewritten, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = _MODIFIED.sub(rb"\g<1>" + modified + rb"\g<2>", data)
            entry = zipfile.ZipInfo(info.filename, date_time=stamp)
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = info.external_attr
            target.writestr(entry, data)
    rewritten.replace(path)


def write_workbook(
    sheets: ExtractionSheets, path: Path, timestamp: datetime | None = None
) -> WorkbookSummary:
    return WorkbookWriter(sheets).write(path, timestamp)

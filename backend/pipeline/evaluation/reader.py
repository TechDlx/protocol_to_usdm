"""Read a legacy USDM workbook back into the review layouts: sheet key -> rows of header -> text.

Both sides of an evaluation go through this reader, the workbook this pipeline wrote and the
reference workbook, so they are compared in the same shape. The reader follows the dialect the
writer produces and the reference workbook uses (D26), tolerating the reference's variations:

- `study`: the key/value block, then the governance dates table whose header row starts with
  `category`;
- `studyDesign`: the key/value block, then the arm-by-epoch element grid (`Epoch/Arms` or
  `Arms/Epochs` in its corner cell), read as one row per arm and epoch;
- timeline sheets, named by `mainTimeline` / `otherTimelines`: Name/Label/Description/Condition in
  columns A-B, heading rows named in column C with one column per timepoint, then the activity
  table after the `Parent Activity` row with `X` marks (a lone `-` is an empty cell);
- every other sheet: headers on row 1, one record per row. `studyDesignInterventions` is read as
  `studyInterventions`.

Content the layouts do not model (sheets, columns and keys the pipeline does not produce, columns
listed with `field=None`) is not returned as rows but counted per location, so an evaluation can
report how much of the reference lies outside the pipeline's scope.
"""

from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from backend.pipeline.workbook.layout import (
    DATES,
    SCHEDULE,
    SHEETS,
    STUDY,
    STUDY_CELLS,
    STUDY_DESIGN,
    TIMELINES,
    TIMEPOINTS,
    SheetKind,
    SheetSpec,
)

Rows = list[dict[str, str]]

SHEET_ALIASES = {"studyInterventions": ("studyInterventions", "studyDesignInterventions")}
GRID_CORNERS = ("epoch/arms", "arms/epochs")
ACTIVITY_HEADER = "parent activity"
TIMELINE_META = {
    "name": "name",
    "label": "label",
    "description": "description",
    "condition": "condition",
}
TIMELINE_KEYS = ("mainTimeline", "otherTimelines")
MARKS = {"x"}
BC_PREFIX = "bc:"
# Layout views that are not sheets of their own; they are filled from the special layouts.
VIRTUAL = {DATES.key, STUDY_CELLS.key, TIMELINES.key, TIMEPOINTS.key, SCHEDULE.key}


@dataclass
class WorkbookTables:
    sheets: dict[str, Rows] = field(default_factory=dict)  # layout key -> rows
    #: Non-empty cells the layouts do not model, by "sheet" or "sheet: column".
    unmodelled: Counter[str] = field(default_factory=Counter)
    missing_sheets: list[str] = field(default_factory=list)  # layout sheets absent from the file


def text(value: Any) -> str:
    """A cell as text: numbers without a spurious .0, dates as yyyy-mm-dd, whitespace trimmed."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _grid(ws: Worksheet) -> list[list[str]]:
    return [[text(v) for v in row] for row in ws.iter_rows(values_only=True)]


def _cell(grid: list[list[str]], row: int, col: int) -> str:
    """0-based access that treats cells beyond the used range as empty."""
    if row >= len(grid) or col >= len(grid[row]):
        return ""
    return grid[row][col]


def _key_value(
    grid: list[list[str]], spec: SheetSpec, tables: WorkbookTables, extra_keys: tuple[str, ...] = ()
) -> tuple[dict[str, str], int]:
    """Read the key/value block at the top of a sheet. Returns the row and the index of the first
    row after the block."""
    columns = {c.header.casefold(): c for c in spec.columns}
    record: dict[str, str] = {}
    row = 0
    while row < len(grid) and _cell(grid, row, 0):
        key, value = _cell(grid, row, 0), _cell(grid, row, 1)
        column = columns.get(key.casefold())
        if column is not None and column.field is not None:
            record[column.header] = value
        elif key in extra_keys:
            record[key] = value
        elif value:
            tables.unmodelled[f"{spec.workbook_sheet}: {key}"] += 1
        row += 1
    return record, row


def _table(
    grid: list[list[str]],
    spec: SheetSpec,
    header_row: int,
    tables: WorkbookTables,
    stop_blank: bool,
) -> Rows:
    """Rows of a table whose headers are on `header_row` (0-based)."""
    headers = [h.strip() for h in (grid[header_row] if header_row < len(grid) else [])]
    columns = {c.header.casefold(): c for c in spec.columns}
    rows: Rows = []
    for raw in grid[header_row + 1 :]:
        if not any(raw):
            if stop_blank:
                break
            continue
        record = {c.header: "" for c in spec.columns if c.field is not None}
        for i, value in enumerate(raw):
            if not value:
                continue
            header = headers[i] if i < len(headers) else ""
            column = columns.get(header.casefold())
            if column is not None and column.field is not None:
                record[column.header] = value
            else:
                tables.unmodelled[f"{spec.workbook_sheet}: {header or '(no header)'}"] += 1
        rows.append(record)
    return rows


def _study(grid: list[list[str]], tables: WorkbookTables) -> None:
    record, row = _key_value(grid, STUDY, tables)
    tables.sheets[STUDY.key] = [record]
    header = next(
        (r for r in range(row, len(grid)) if _cell(grid, r, 0).casefold() == "category"), None
    )
    tables.sheets[DATES.key] = (
        _table(grid, DATES, header, tables, stop_blank=True) if header is not None else []
    )


def _study_design(grid: list[list[str]], tables: WorkbookTables) -> list[tuple[str, bool]]:
    """Read the design block and grid. Returns the timeline sheet names, main first."""
    record, row = _key_value(grid, STUDY_DESIGN, tables, extra_keys=TIMELINE_KEYS)
    timelines: list[tuple[str, bool]] = []
    main = record.pop("mainTimeline", "")
    others = record.pop("otherTimelines", "")
    if main:
        timelines.append((main, True))
    timelines += [(name.strip(), False) for name in others.split(",") if name.strip()]
    tables.sheets[STUDY_DESIGN.key] = [record]

    cells: Rows = []
    corner = next(
        (r for r in range(row, len(grid)) if _cell(grid, r, 0).casefold() in GRID_CORNERS), None
    )
    if corner is not None:
        epochs = grid[corner][1:]
        for arm_row in grid[corner + 1 :]:
            arm = arm_row[0] if arm_row else ""
            if not arm:
                break
            for i, epoch in enumerate(epochs, start=1):
                element = arm_row[i] if i < len(arm_row) else ""
                if epoch and element:
                    cells.append({"arm": arm, "epoch": epoch, "elements": element})
    tables.sheets[STUDY_CELLS.key] = cells
    return timelines


def _timeline(grid: list[list[str]], sheet: str, main: bool, tables: WorkbookTables) -> None:
    def blank_dash(value: str) -> str:
        return "" if value == "-" else value

    activity_row = next(
        (r for r in range(len(grid)) if _cell(grid, r, 0).casefold() == ACTIVITY_HEADER), len(grid)
    )
    meta: dict[str, str] = {}
    headings: dict[str, list[str]] = {}
    for r in range(activity_row):
        key = _cell(grid, r, 0).casefold()
        if key in TIMELINE_META:
            meta[TIMELINE_META[key]] = _cell(grid, r, 1)
        heading = _cell(grid, r, 2).casefold()
        if heading:
            headings[heading] = [blank_dash(v) for v in grid[r][3:]]
    name = meta.get("name", "") or sheet
    tables.sheets[TIMELINES.key].append(
        {
            "name": name,
            "sheet": sheet,
            "mainTimeline": "Y" if main else "N",
            "description": meta.get("description", ""),
            "label": meta.get("label", ""),
            "condition": meta.get("condition", ""),
        }
    )

    names = headings.get("name", [])
    count = max((i + 1 for i, n in enumerate(names) if n), default=0)
    fields = [c.header for c in TIMEPOINTS.columns if c.header != "timeline"]
    for i in range(count):
        row = {"timeline": name}
        for header in fields:
            values = headings.get(header.casefold(), [])
            row[header] = values[i] if i < len(values) else ""
        tables.sheets[TIMEPOINTS.key].append(row)

    for raw in grid[activity_row + 1 :]:
        activity = _cell([raw], 0, 1) or _cell([raw], 0, 0)
        if not activity:
            continue
        concepts, others = [], 0
        for item in _cell([raw], 0, 2).split(","):
            item = item.strip()
            if item.casefold().startswith(BC_PREFIX):
                concepts.append(item[len(BC_PREFIX) :].strip())
            elif item:
                others += 1
        if others:
            tables.unmodelled[f"{sheet}: procedures and sub-timelines"] += others
        marks = [
            names[i]
            for i in range(count)
            if blank_dash(_cell([raw], 0, 3 + i)).casefold() in MARKS and names[i]
        ]
        tables.sheets[SCHEDULE.key].append(
            {
                "timeline": name,
                "activity": activity,
                "biomedicalConcepts": ", ".join(concepts),
                "scheduledAt": ", ".join(marks),
            }
        )


def read_workbook(path: Path) -> WorkbookTables:
    book = load_workbook(path, data_only=True, read_only=False)
    tables = WorkbookTables()
    grids = {name: _grid(book[name]) for name in book.sheetnames}
    used: set[str] = set()

    if STUDY.workbook_sheet in grids:
        _study(grids[STUDY.workbook_sheet], tables)
        used.add(STUDY.workbook_sheet)
    else:
        tables.missing_sheets.append(STUDY.workbook_sheet)
    timelines: list[tuple[str, bool]] = []
    if STUDY_DESIGN.workbook_sheet in grids:
        timelines = _study_design(grids[STUDY_DESIGN.workbook_sheet], tables)
        used.add(STUDY_DESIGN.workbook_sheet)
    else:
        tables.missing_sheets.append(STUDY_DESIGN.workbook_sheet)

    tables.sheets.update({TIMELINES.key: [], TIMEPOINTS.key: [], SCHEDULE.key: []})
    for sheet, main in timelines:
        if sheet in grids:
            _timeline(grids[sheet], sheet, main, tables)
            used.add(sheet)

    for spec in SHEETS.values():
        if spec.kind != SheetKind.TABLE or spec.key in VIRTUAL:
            continue
        name = next(
            (
                n
                for n in SHEET_ALIASES.get(spec.workbook_sheet, (spec.workbook_sheet,))
                if n in grids
            ),
            None,
        )
        if name is None:
            tables.missing_sheets.append(spec.workbook_sheet)
            continue
        used.add(name)
        tables.sheets[spec.key] = _table(
            grids[name], spec, spec.first_row - 2, tables, stop_blank=False
        )

    for name, grid in grids.items():
        if name not in used:
            filled = sum(1 for row in grid for value in row if value)
            if filled:
                tables.unmodelled[name] += filled
    book.close()
    return tables
